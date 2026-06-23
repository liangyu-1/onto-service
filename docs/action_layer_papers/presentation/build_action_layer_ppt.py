from pathlib import Path

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("action_layer_research_plan_zh.pptx")

W = Inches(13.333)
H = Inches(7.5)
FONT = "PingFang SC"

NAVY = RGBColor(22, 42, 63)
BLUE = RGBColor(42, 103, 158)
TEAL = RGBColor(20, 126, 125)
GREEN = RGBColor(52, 132, 91)
AMBER = RGBColor(210, 139, 35)
RED = RGBColor(180, 67, 62)
INK = RGBColor(36, 44, 52)
MUTED = RGBColor(91, 104, 116)
LINE = RGBColor(211, 218, 224)
PALE_BLUE = RGBColor(235, 243, 249)
PALE_TEAL = RGBColor(232, 246, 244)
PALE_AMBER = RGBColor(251, 245, 231)
PALE_RED = RGBColor(250, 237, 236)
WHITE = RGBColor(255, 255, 255)
BG = RGBColor(248, 250, 251)


def set_run(run, size=20, bold=False, color=INK, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_text(slide, x, y, w, h, text, size=20, bold=False, color=INK,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    return box


def add_rich_text(slide, x, y, w, h, lines, size=18, color=INK,
                  bullet=False, spacing=8):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.05)
    frame.margin_top = Inches(0.03)
    frame.margin_bottom = Inches(0.03)
    for idx, item in enumerate(lines):
        if isinstance(item, tuple):
            text, is_bold, item_color = item
        else:
            text, is_bold, item_color = item, False, color
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = text
        p.space_after = Pt(spacing)
        p.level = 0
        if bullet:
            p.text = "• " + p.text
        for run in p.runs:
            set_run(run, size, is_bold, item_color)
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def line(slide, x1, y1, x2, y2, color=LINE, width=1.5, arrow=False):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    if arrow:
        shape.line.end_arrowhead = True
    return shape


def add_header(slide, title, subtitle=None, section=None):
    if section:
        add_text(slide, 0.65, 0.25, 2.2, 0.3, section.upper(), 10, True, TEAL)
    add_text(slide, 0.65, 0.55, 12.0, 0.55, title, 27, True, NAVY)
    if subtitle:
        add_text(slide, 0.68, 1.12, 11.8, 0.38, subtitle, 13, False, MUTED)
    line(slide, 0.65, 1.55, 12.65, 1.55, LINE, 1)


def add_footer(slide, number):
    add_text(slide, 0.68, 7.12, 5.0, 0.2, "Ontology Action Layer Research Plan", 8, False, MUTED)
    add_text(slide, 12.15, 7.08, 0.5, 0.25, str(number), 9, True, MUTED, PP_ALIGN.RIGHT)


def add_badge(slide, x, y, w, text, fill, color=WHITE):
    shape = rect(slide, x, y, w, 0.34, fill, fill)
    shape.line.fill.background()
    add_text(slide, x, y + 0.01, w, 0.28, text, 10, True, color, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def add_card(slide, x, y, w, h, title, body, accent=TEAL, body_size=16):
    rect(slide, x, y, w, h, WHITE, LINE)
    rect(slide, x, y, 0.08, h, accent, accent, False)
    add_text(slide, x + 0.28, y + 0.18, w - 0.48, 0.38, title, 17, True, NAVY)
    add_rich_text(slide, x + 0.25, y + 0.66, w - 0.45, h - 0.82, body, body_size, MUTED, True, 7)


def new_slide(prs, title=None, subtitle=None, section=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    if title:
        add_header(slide, title, subtitle, section)
    add_footer(slide, len(prs.slides))
    return slide


def stage_box(slide, x, y, w, title, subtitle, color):
    rect(slide, x, y, w, 1.15, WHITE, color)
    add_badge(slide, x + 0.18, y + 0.14, 0.62, title, color)
    add_text(slide, x + 0.18, y + 0.58, w - 0.36, 0.42, subtitle, 14, True, NAVY, PP_ALIGN.CENTER)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "面向本体动作层的研究：构建、检索与智能体执行"
    prs.core_properties.subject = "三篇论文研究计划导师汇报"
    prs.core_properties.author = "Yuliang"

    # 1 Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    rect(slide, 0, 0, 0.18, 7.5, TEAL, TEAL, False)
    add_badge(slide, 0.85, 0.72, 1.65, "RESEARCH PLAN", TEAL)
    add_text(slide, 0.85, 1.42, 11.4, 1.2, "面向本体动作层的研究", 36, True, WHITE)
    add_text(slide, 0.87, 2.6, 11.0, 0.65, "构建 · 检索 · 智能体理解与规划", 24, False, RGBColor(198, 216, 228))
    line(slide, 0.87, 3.55, 5.7, 3.55, TEAL, 3)
    add_text(slide, 0.87, 4.0, 10.8, 0.8,
             "核心目标：让本体从“描述世界”扩展到“支持可执行动作”", 20, True, WHITE)
    add_text(slide, 0.87, 6.55, 7.5, 0.32, "导师汇报 · 2026年6月", 12, False, RGBColor(170, 194, 210))
    add_text(slide, 11.7, 6.48, 0.75, 0.42, "01", 20, True, TEAL, PP_ALIGN.RIGHT)

    # 2 Motivation
    slide = new_slide(prs, "研究缺口：本体描述了对象，但没有充分表达行动", section="MOTIVATION")
    add_card(slide, 0.7, 1.85, 3.65, 3.65, "传统本体", [
        "对象、属性与关系",
        "类型层次和约束",
        "适合查询“是什么”",
    ], BLUE, 17)
    add_text(slide, 4.55, 3.1, 0.8, 0.5, "≠", 32, True, RED, PP_ALIGN.CENTER)
    add_card(slide, 5.45, 1.85, 3.65, 3.65, "工具 / API Schema", [
        "函数名和 JSON 参数",
        "规定调用语法",
        "不完整表达业务状态",
    ], AMBER, 17)
    add_text(slide, 9.3, 3.1, 0.8, 0.5, "→", 30, True, TEAL, PP_ALIGN.CENTER)
    add_card(slide, 10.1, 1.85, 2.55, 3.65, "Action Layer", [
        "能做什么",
        "何时能做",
        "作用于谁",
        "执行后怎样",
    ], TEAL, 16)
    rect(slide, 1.75, 5.85, 9.9, 0.66, PALE_RED, RGBColor(232, 190, 187))
    add_text(slide, 1.95, 6.02, 9.5, 0.3,
             "研究对象不是“再做一个知识图谱”，而是操作型本体中的动作语义。", 16, True, RED, PP_ALIGN.CENTER)

    # 3 Definition
    slide = new_slide(prs, "Action Layer 到底是什么？", "动作不是一个动词标签，而是带状态语义的可执行本体实体", "FOUNDATION")
    cx, cy = 6.65, 3.95
    rect(slide, 5.1, 3.15, 3.1, 1.5, NAVY, NAVY)
    add_text(slide, 5.25, 3.48, 2.8, 0.55, "Ontology Action", 25, True, WHITE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    nodes = [
        (0.8, 2.0, 2.3, "Target Object", "动作作用于哪个对象", BLUE),
        (0.8, 4.7, 2.3, "Evidence", "来自 SOP 的证据与来源", TEAL),
        (3.45, 1.72, 2.1, "Actor / Role", "谁有权执行", GREEN),
        (3.45, 5.0, 2.1, "Parameters", "参数及语义角色", AMBER),
        (7.75, 1.72, 2.1, "Preconditions", "当前状态是否允许", RED),
        (7.75, 5.0, 2.1, "Effects", "状态如何迁移", BLUE),
        (10.35, 2.0, 2.2, "Constraints", "策略、确认与安全规则", AMBER),
        (10.35, 4.7, 2.2, "Control Flow", "前驱、后继与异常分支", TEAL),
    ]
    for x, y, w, title, body, color in nodes:
        rect(slide, x, y, w, 1.05, WHITE, color)
        add_text(slide, x + 0.1, y + 0.13, w - 0.2, 0.27, title, 14, True, color, PP_ALIGN.CENTER)
        add_text(slide, x + 0.1, y + 0.49, w - 0.2, 0.35, body, 11, False, MUTED, PP_ALIGN.CENTER)
        line(slide, x + w / 2, y + 0.52, cx, cy, color, 1.2)

    # 4 Thesis map
    slide = new_slide(prs, "三篇论文构成一条 Action Layer 闭环", "逻辑依赖顺序不等于实际写作顺序", "RESEARCH MAP")
    stage_box(slide, 0.7, 2.15, 3.2, "Paper 2", "从文档构建动作", TEAL)
    stage_box(slide, 5.05, 2.15, 3.2, "Paper 1", "从动作库检索候选", BLUE)
    stage_box(slide, 9.4, 2.15, 3.2, "Paper 3", "Agent 理解、约束与修复", AMBER)
    line(slide, 3.92, 2.73, 5.0, 2.73, TEAL, 3, True)
    line(slide, 8.27, 2.73, 9.35, 2.73, BLUE, 3, True)
    add_text(slide, 1.05, 3.55, 2.5, 0.65, "Document + Ontology\n→ Action IR", 16, True, TEAL, PP_ALIGN.CENTER)
    add_text(slide, 5.38, 3.55, 2.5, 0.65, "Query + State\n→ Ranked Actions", 16, True, BLUE, PP_ALIGN.CENTER)
    add_text(slide, 9.72, 3.55, 2.5, 0.65, "Goal + Runtime State\n→ Safe Execution", 16, True, AMBER, PP_ALIGN.CENTER)
    rect(slide, 1.2, 5.05, 10.95, 1.1, PALE_BLUE, RGBColor(188, 210, 226))
    add_text(slide, 1.45, 5.25, 2.35, 0.28, "实际推进顺序", 14, True, NAVY)
    add_text(slide, 3.65, 5.17, 7.8, 0.55,
             "Paper 3（已有验证） → Paper 2（补齐动作来源） → Paper 1（规模化后检索）",
             16, True, BLUE, PP_ALIGN.CENTER)

    # 5 Shared IR
    slide = new_slide(prs, "统一接口：Action IR 贯通三篇论文", "避免每篇论文各自定义一套动作表示", "FOUNDATION")
    fields = [
        ("Identity", "action_id\naction_type\naction_kind", BLUE),
        ("Grounding", "actor\ntarget_object\nrole_bindings", TEAL),
        ("Semantics", "preconditions\neffects\nconstraints", RED),
        ("Evidence", "evidence span\nprovenance\nconfidence", GREEN),
        ("Workflow", "control_flow\ngrounding_mode\nruntime projection", AMBER),
    ]
    for i, (title, body, color) in enumerate(fields):
        x = 0.68 + i * 2.52
        rect(slide, x, 2.0, 2.25, 2.55, WHITE, color)
        add_badge(slide, x + 0.32, 2.22, 1.61, title, color)
        add_text(slide, x + 0.2, 2.95, 1.85, 1.15, body, 15, False, INK, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    rect(slide, 1.35, 5.18, 10.6, 0.82, NAVY, NAVY)
    add_text(slide, 1.55, 5.38, 10.2, 0.32,
             "Paper 2 生成与验证 → Paper 1 建索引与排序 → Paper 3 投影为 ActionBank 并执行",
             17, True, WHITE, PP_ALIGN.CENTER)

    # 6 Paper 2 problem
    slide = new_slide(prs, "Paper 2：从非结构化程序文档构建动作层", "区别于普通事件抽取：输出必须可 grounding、可验证、可执行", "PAPER 2")
    add_card(slide, 0.72, 1.88, 3.65, 3.95, "输入", [
        "SOP / 维护手册",
        "客服脚本 / 工作指令",
        "已有对象与类型本体",
    ], TEAL, 17)
    add_text(slide, 4.48, 3.28, 0.65, 0.45, "+", 28, True, NAVY, PP_ALIGN.CENTER)
    add_card(slide, 5.15, 1.88, 3.2, 3.95, "核心难点", [
        "结构：标题、表格、编号",
        "语义：条件、异常、效果",
        "对齐：动作与对象类型",
        "证据：禁止无来源生成",
    ], RED, 16)
    add_text(slide, 8.42, 3.28, 0.65, 0.45, "→", 28, True, NAVY, PP_ALIGN.CENTER)
    add_card(slide, 9.08, 1.88, 3.55, 3.95, "输出", [
        "Grounded Action IR",
        "Procedure Graph",
        "验证结果与证据链",
    ], BLUE, 17)
    add_text(slide, 1.35, 6.25, 10.65, 0.4,
             "主张：本体约束应提升结构完整性与语义一致性，而不只是动作短语 F1。", 16, True, TEAL, PP_ALIGN.CENTER)

    # 7 Paper 2 method
    slide = new_slide(prs, "Paper 2 方法：LLM 抽取 + 本体约束 + 证据验证", section="PAPER 2")
    steps = [
        ("1", "结构感知解析", "版面、标题、表格、步骤号", TEAL),
        ("2", "程序单元识别", "动作、条件、异常、安全规则", BLUE),
        ("3", "Action IR 生成", "LLM 输出受控结构与证据", AMBER),
        ("4", "Ontology Grounding", "候选检索、类型与角色对齐", GREEN),
        ("5", "程序图与验证", "流程边、约束、可投影性检查", RED),
    ]
    for i, (num, title, body, color) in enumerate(steps):
        x = 0.65 + i * 2.53
        rect(slide, x, 2.12, 2.22, 3.15, WHITE, color)
        add_badge(slide, x + 0.77, 2.35, 0.68, num, color)
        add_text(slide, x + 0.13, 3.08, 1.96, 0.45, title, 16, True, NAVY, PP_ALIGN.CENTER)
        add_text(slide, x + 0.18, 3.82, 1.86, 0.8, body, 14, False, MUTED, PP_ALIGN.CENTER)
        if i < 4:
            line(slide, x + 2.24, 3.68, x + 2.48, 3.68, color, 2.2, True)
    rect(slide, 1.25, 5.75, 10.85, 0.62, PALE_AMBER, RGBColor(230, 208, 162))
    add_text(slide, 1.45, 5.9, 10.45, 0.28,
             "关键边界：不从文档“自由生成本体”；优先 grounding 到已有本体，无法对齐的字段显式 unresolved。",
             14, True, AMBER, PP_ALIGN.CENTER)

    # 8 Paper 2 evaluation
    slide = new_slide(prs, "Paper 2 实验：必须同时评估抽取、grounding 与流程", section="PAPER 2")
    add_card(slide, 0.72, 1.85, 3.65, 4.55, "数据集", [
        "PET：活动、参与者、控制流",
        "MyFixit：维修动作、工具、零件",
        "MSPT：科学步骤与参数",
        "OMIn：维护实体与事件",
        "人工审核的完整 Action IR 子集",
    ], TEAL, 15)
    add_card(slide, 4.83, 1.85, 3.65, 4.55, "对照方法", [
        "规则与序列标注模型",
        "LLM-only extraction",
        "LLM + ontology prompt",
        "Ours：候选检索 + symbolic validation",
    ], BLUE, 16)
    add_card(slide, 8.94, 1.85, 3.65, 4.55, "评价指标", [
        "Action / Evidence F1",
        "Type / Target Grounding Accuracy",
        "Parameter F1",
        "Control-flow Edge F1",
        "Validation Pass Rate",
        "下游 retrieval / planning utility",
    ], AMBER, 15)

    # 9 Paper 1 premise
    slide = new_slide(prs, "Paper 1：Action 为什么需要检索？", "一对一绑定时检索价值有限；研究问题成立于多动作、动态状态和自然语言任务", "PAPER 1")
    add_card(slide, 0.72, 1.95, 4.0, 3.65, "不成立的简单场景", [
        "一个对象只绑定一个动作",
        "动作规模很小",
        "调用入口已经确定",
        "当前状态不影响选择",
    ], RED, 17)
    add_text(slide, 4.86, 3.22, 0.8, 0.52, "VS", 24, True, MUTED, PP_ALIGN.CENTER)
    add_card(slide, 5.75, 1.95, 6.85, 3.65, "真正的检索场景", [
        "同一对象对应查询、修改、审批、诊断、修复等多个动作",
        "相似动作由对象状态、权限、参数和前置条件区分",
        "动作来自多份 SOP、API 和业务域，规模持续扩大",
        "用户只描述目标，不知道具体动作名称",
    ], TEAL, 17)
    rect(slide, 1.8, 5.96, 9.7, 0.58, NAVY, NAVY)
    add_text(slide, 2.0, 6.1, 9.3, 0.28,
             "任务：Query + Runtime State + Ontology → Ranked Executable Actions",
             16, True, WHITE, PP_ALIGN.CENTER)

    # 10 Paper 1 architecture
    slide = new_slide(prs, "Paper 1 方法：四路召回与可执行性重排", section="PAPER 1")
    add_text(slide, 0.82, 2.02, 2.0, 0.45, "自然语言任务 + 状态", 17, True, NAVY, PP_ALIGN.CENTER)
    line(slide, 2.86, 2.25, 3.45, 2.25, NAVY, 2.5, True)
    channels = [
        (3.55, 1.72, "Lexical", "动作名、对象名\n参数与领域术语", BLUE),
        (5.55, 1.72, "Dense", "语义改写\n文档证据", TEAL),
        (7.55, 1.72, "Graph", "对象—状态—动作\n流程邻域", GREEN),
        (9.55, 1.72, "Constraint", "前置条件\n参数与状态迁移", RED),
    ]
    for x, y, title, body, color in channels:
        rect(slide, x, y, 1.7, 1.55, WHITE, color)
        add_badge(slide, x + 0.28, y + 0.18, 1.14, title, color)
        add_text(slide, x + 0.15, y + 0.72, 1.4, 0.62, body, 12, False, MUTED, PP_ALIGN.CENTER)
        line(slide, x + 0.85, y + 1.58, 7.4, 4.05, color, 1.4)
    rect(slide, 4.72, 3.95, 5.35, 1.3, NAVY, NAVY)
    add_text(slide, 4.9, 4.12, 5.0, 0.35, "Executability-Aware Reranker", 20, True, WHITE, PP_ALIGN.CENTER)
    add_text(slide, 4.9, 4.55, 5.0, 0.35, "Relevance + Grounding + State + Evidence", 13, False, RGBColor(199, 218, 230), PP_ALIGN.CENTER)
    line(slide, 10.1, 4.6, 11.0, 4.6, NAVY, 2.5, True)
    rect(slide, 11.08, 3.95, 1.55, 1.3, PALE_TEAL, TEAL)
    add_text(slide, 11.22, 4.18, 1.27, 0.72, "Top-k\n可执行动作", 15, True, TEAL, PP_ALIGN.CENTER)
    add_text(slide, 1.0, 5.95, 11.4, 0.42,
             "区别：不是先按文本相似度检索、最后再过滤；可执行性直接参与排序。",
             16, True, BLUE, PP_ALIGN.CENTER)

    # 11 Paper 1 evaluation
    slide = new_slide(prs, "Paper 1 实验：证明“相关”不等于“可执行”", section="PAPER 1")
    add_card(slide, 0.72, 1.9, 3.65, 4.25, "Benchmark 构造", [
        "从 Paper 2 的 Action IR 生成查询",
        "状态反事实：Pending / Delivered / Cancelled",
        "角色扰动：错误 item、payment、target",
        "相似动作难负例",
    ], TEAL, 16)
    add_card(slide, 4.83, 1.9, 3.65, 4.25, "Baseline", [
        "BM25 / Dense / Hybrid",
        "GraphRAG / KG-RAG",
        "Ontology concept index",
        "Tool retrieval / LLM-only",
    ], BLUE, 16)
    add_card(slide, 8.94, 1.9, 3.65, 4.25, "核心证据", [
        "Recall@k、MRR、nDCG",
        "Executable Hit@k",
        "Constraint Satisfaction",
        "State-sensitive Ranking Accuracy",
        "Path / Evidence Faithfulness",
    ], AMBER, 15)
    add_text(slide, 1.4, 6.45, 10.6, 0.28,
             "风险：公开数据缺少完整 precondition / effect 标注，必须建设人工审核子集。",
             14, True, RED, PP_ALIGN.CENTER)

    # 12 Paper 3 method
    slide = new_slide(prs, "Paper 3：本体约束的 Agent 动作理解与规划", "核心不是通用 planner，而是 LLM 提案与工具执行之间的动作语义层", "PAPER 3")
    flow = [
        ("Goal + History", "用户目标与已观察轨迹", BLUE),
        ("LLM Proposal", "候选 action + arguments", AMBER),
        ("Ontology State", "对象、状态、角色绑定", TEAL),
        ("Admissibility", "类型感知的可执行性检查", RED),
        ("Execute / Repair", "执行或结构化修复", GREEN),
    ]
    for i, (title, body, color) in enumerate(flow):
        x = 0.62 + i * 2.55
        rect(slide, x, 2.05, 2.2, 1.5, WHITE, color)
        add_text(slide, x + 0.1, 2.28, 2.0, 0.35, title, 16, True, color, PP_ALIGN.CENTER)
        add_text(slide, x + 0.15, 2.82, 1.9, 0.42, body, 12, False, MUTED, PP_ALIGN.CENTER)
        if i < 4:
            line(slide, x + 2.22, 2.8, x + 2.48, 2.8, color, 2.2, True)
    add_card(slide, 1.0, 4.35, 5.25, 1.6, "Epistemic Action：弱 Grounding", [
        "查询、检查、诊断用于获取状态",
        "不能要求目标对象已经被观察",
    ], BLUE, 15)
    add_card(slide, 7.05, 4.35, 5.25, 1.6, "Mutating Action：强 Grounding", [
        "修改、取消、退款会改变状态",
        "必须检查对象、角色、状态、策略和确认",
    ], RED, 15)

    # 13 Results
    slide = new_slide(prs, "Paper 3 初步结果：积极，但尚不能过度解释", "主分析：3 个配额完整 trial，341 个严格配对样本", "PAPER 3")
    chart_data = ChartData()
    chart_data.categories = ["Schema-Only", "Typed + Repair"]
    chart_data.add_series("Success rate", (44.9, 55.7))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.85), Inches(1.9), Inches(5.6), Inches(4.35), chart_data
    ).chart
    chart.has_legend = False
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = 70
    chart.value_axis.major_unit = 10
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.tick_labels.font.name = FONT
    chart.value_axis.tick_labels.font.size = Pt(11)
    chart.category_axis.tick_labels.font.name = FONT
    chart.category_axis.tick_labels.font.size = Pt(12)
    chart.plots[0].has_data_labels = True
    chart.plots[0].data_labels.number_format = '0.0"%"'
    chart.plots[0].data_labels.font.size = Pt(14)
    chart.plots[0].data_labels.font.bold = True
    chart.plots[0].series[0].format.fill.solid()
    chart.plots[0].series[0].format.fill.fore_color.rgb = TEAL
    add_card(slide, 6.9, 1.9, 5.7, 2.05, "统计结论", [
        "+10.9 个百分点",
        "paired McNemar p = 0.0029",
        "task-level sign test p = 0.109（未显著）",
    ], TEAL, 16)
    add_card(slide, 6.9, 4.22, 5.7, 2.05, "机制与代价", [
        "Transfer calls：323 → 62",
        "工具调用：8.30 → 7.16 / simulation",
        "平均延迟：190.6s → 260.3s",
    ], AMBER, 16)

    # 14 Critical limitations
    slide = new_slide(prs, "必须向导师讲清楚的限制", "这些限制决定当前结论只能称为初步证据", "CRITICAL REVIEW")
    add_card(slide, 0.72, 1.86, 3.75, 4.5, "实验完整性", [
        "计划 8 trials，后半段因模型配额中断",
        "主结果只使用双方完整的 trials 0–2",
        "不能把 infrastructure error 当任务失败",
    ], RED, 16)
    add_card(slide, 4.8, 1.86, 3.75, 4.5, "方法缺陷", [
        "发现 confirmation binding 缺陷",
        "修复版仅通过 symbolic smoke test",
        "尚无修复后的端到端结果",
    ], AMBER, 16)
    add_card(slide, 8.88, 1.86, 3.75, 4.5, "论文主张边界", [
        "simulation-level 显著",
        "task-level sign test 尚未显著",
        "不能宣称已解决通用 planning",
        "当前证明的是动作语义层的可行性",
    ], BLUE, 15)

    # 15 Relationship and validation
    slide = new_slide(prs, "三篇论文分别要证明什么？", "每篇论文只承担一个主要因果主张", "SYNTHESIS")
    rows = [
        ("Paper 2", "Ontology constraint", "Action IR 更完整、更一致、可追溯", TEAL),
        ("Paper 1", "Executability signal", "同等召回下，可执行动作排名更高", BLUE),
        ("Paper 3", "Type-aware grounding", "保留读动作，约束写动作，并改善执行", AMBER),
    ]
    add_text(slide, 0.85, 1.9, 2.2, 0.35, "论文", 14, True, MUTED)
    add_text(slide, 3.15, 1.9, 3.0, 0.35, "被检验的机制", 14, True, MUTED)
    add_text(slide, 6.5, 1.9, 5.4, 0.35, "可证伪的结果", 14, True, MUTED)
    line(slide, 0.8, 2.35, 12.4, 2.35, LINE, 1.2)
    for i, (paper, mechanism, claim, color) in enumerate(rows):
        y = 2.65 + i * 1.15
        add_badge(slide, 0.85, y, 1.35, paper, color)
        add_text(slide, 3.12, y + 0.02, 2.85, 0.48, mechanism, 17, True, color)
        add_text(slide, 6.48, y + 0.02, 5.65, 0.58, claim, 17, False, INK)
        line(slide, 0.8, y + 0.82, 12.4, y + 0.82, LINE, 0.8)
    rect(slide, 1.35, 6.2, 10.55, 0.56, PALE_RED, RGBColor(232, 190, 187))
    add_text(slide, 1.55, 6.33, 10.15, 0.27,
             "若三篇论文都只比较最终准确率，则无法证明“本体动作层”是有效机制。",
             14, True, RED, PP_ALIGN.CENTER)

    # 16 Roadmap
    slide = new_slide(prs, "建议推进路线：先闭合证据，再扩展研究范围", section="ROADMAP")
    phases = [
        ("01", "Paper 3 重跑", "修复 confirmation\n完整 paired run\ncomponent ablation", RED),
        ("02", "统一标注集", "Action IR guideline\n人工审核子集\n跨数据集映射", TEAL),
        ("03", "Paper 2 抽取", "训练与验证\n误差分析\n生成 ActionBank", BLUE),
        ("04", "Paper 1 检索", "状态反事实\n可执行性重排\n跨域泛化", AMBER),
    ]
    for i, (num, title, body, color) in enumerate(phases):
        x = 0.72 + i * 3.15
        rect(slide, x, 2.05, 2.75, 3.65, WHITE, color)
        add_badge(slide, x + 0.2, 2.27, 0.68, num, color)
        add_text(slide, x + 0.2, 2.95, 2.35, 0.42, title, 18, True, NAVY, PP_ALIGN.CENTER)
        add_text(slide, x + 0.2, 3.72, 2.35, 1.05, body, 15, False, MUTED, PP_ALIGN.CENTER)
        if i < 3:
            line(slide, x + 2.78, 3.88, x + 3.08, 3.88, color, 2.4, True)
    add_text(slide, 1.2, 6.18, 10.9, 0.38,
             "跨域顺序：Retail → Airline → Maintenance / Industrial SOP",
             17, True, TEAL, PP_ALIGN.CENTER)

    # 17 Questions
    slide = new_slide(prs, "希望导师重点判断的问题", section="DISCUSSION")
    questions = [
        ("1", "Action Layer 是否足以作为大论文的统一研究对象？", TEAL),
        ("2", "Paper 1 的 retrieval 与 Paper 3 的 planning 边界是否清晰？", BLUE),
        ("3", "Paper 2 是否应先聚焦一种文档域，而不是同时覆盖所有 SOP？", AMBER),
        ("4", "Paper 3 当前证据是否值得继续投入完整重跑与消融？", RED),
    ]
    for i, (num, text, color) in enumerate(questions):
        y = 1.85 + i * 1.18
        rect(slide, 1.05, y, 11.25, 0.88, WHITE, color)
        add_badge(slide, 1.28, y + 0.25, 0.58, num, color)
        add_text(slide, 2.18, y + 0.18, 9.65, 0.48, text, 18, True, NAVY, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, 4.35, 6.68, 4.65, 0.38, "谢谢，请批评指正", 18, True, TEAL, PP_ALIGN.CENTER)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
